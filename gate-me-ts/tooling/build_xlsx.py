import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule
from register import SUBJECTS, MOCK_SLOTS

P = json.load(open("plan.json"))
topics, tts, mocks, rel = P["topics"], P["tts"], P["mock_content"], P["releases"]
subj = {s: n for s, n, _ in SUBJECTS}
order = [s for s, _, _ in SUBJECTS]
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3864")
B = Font(name=F); BB = Font(name=F, bold=True)
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
IMPF = {"A": PatternFill("solid", fgColor="F8CBAD"), "B": PatternFill("solid", fgColor="FFE699"), "C": PatternFill("solid", fgColor="E2EFDA")}
HIT = PatternFill("solid", fgColor="9BC2E6")

def header(ws, row, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v); c.font = H; c.fill = HF; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w

def cell(ws, r, c, v, font=B, fill=None, center=False):
    x = ws.cell(r, c, v); x.font = font; x.border = BOX
    if fill: x.fill = fill
    x.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=not center)
    return x

wb = Workbook()
ov = wb.active; ov.title = "Overview"
reg = wb.create_sheet("Topic Register")
tt = wb.create_sheet("Topic Tests")
mk = wb.create_sheet("Mock Blueprint")
cal = wb.create_sheet("Release Calendar")

# ---------- Topic Register ----------
cols = ["Topic ID", "Subject", "Topic test", "Topic", "Importance", "Q in topic test", "Q in subject test"] + \
       [f"FLT{i+1:02d}" for i in range(15)] + ["Mocks it appears in", "Revisions after topic test", "Revision rule met?"]
header(reg, 1, cols, [9, 24, 9, 52, 11, 9, 9] + [6.5]*15 + [10, 11, 11])
reg.row_dimensions[1].height = 45
for r, t in enumerate(topics, 2):
    cell(reg, r, 1, t["id"], center=True); cell(reg, r, 2, t["subject_name"]); cell(reg, r, 3, t["tt"], center=True)
    cell(reg, r, 4, t["name"]); cell(reg, r, 5, t["imp"], BB, IMPF[t["imp"]], True)
    cell(reg, r, 6, t["tt_q"], center=True); cell(reg, r, 7, t["st_q"], center=True)
    for m in range(15):
        n = t["mocks"].count(m)
        cell(reg, r, 8 + m, n if n else None, center=True, fill=HIT if n else None)
    cell(reg, r, 23, f"=COUNTIF(H{r}:V{r},\">0\")", center=True)
    cell(reg, r, 24, f"=IF(G{r}>0,1,0)+W{r}", BB, center=True)
    cell(reg, r, 25, f'=IF(AND(X{r}>=2,W{r}>=IF(E{r}="A",4,IF(E{r}="B",2,1))),"Yes","NO")', center=True)
last = len(topics) + 1
reg.freeze_panes = "E2"; reg.auto_filter.ref = f"A1:Y{last}"
reg.conditional_formatting.add(f"Y2:Y{last}", CellIsRule(operator="equal", formula=['"NO"'], fill=PatternFill("solid", fgColor="FF7C80")))

# ---------- Topic Tests ----------
header(tt, 1, ["Test", "Subject", "Topic-test title", "Topics covered (questions)", "Questions", "Marks", "Minutes", "Release"],
       [7, 24, 34, 90, 9, 7, 8, 11])
for r, x in enumerate(sorted(tts, key=lambda x: x["no"]), 2):
    ts = [t for t in topics if t["tt"] == x["code"]]
    cell(tt, r, 1, f"TT{x['no']:02d}", BB, center=True); cell(tt, r, 2, subj[x["subject"]]); cell(tt, r, 3, x["title"])
    cell(tt, r, 4, "; ".join(f"{t['name']} ({t['tt_q']})" for t in ts))
    cell(tt, r, 5, sum(t["tt_q"] for t in ts), center=True); cell(tt, r, 6, 50, center=True); cell(tt, r, 7, 90, center=True)
    cell(tt, r, 8, x["date"], center=True)
tt.freeze_panes = "B2"

# ---------- Mock Blueprint ----------
id2 = {t["id"]: t for t in topics}
header(mk, 1, ["Subject", "1-mark Q", "2-mark Q", "Marks"] + [f"FLT{i+1:02d} topics" for i in range(15)],
       [24, 8, 8, 7] + [38]*15)
for r, s in enumerate(order, 2):
    a, b = MOCK_SLOTS[s]
    cell(mk, r, 1, subj[s], BB); cell(mk, r, 2, a, center=True); cell(mk, r, 3, b, center=True)
    cell(mk, r, 4, f"=B{r}+2*C{r}", center=True)
    for m in range(15):
        cell(mk, r, 5 + m, "\n".join(id2[i]["name"] for i in mocks[m].get(s, [])))
tr = len(order) + 2
cell(mk, tr, 1, "TOTAL", BB)
for c, col in ((2, "B"), (3, "C"), (4, "D")):
    cell(mk, tr, c, f"=SUM({col}2:{col}{tr-1})", BB, center=True)
cell(mk, tr + 1, 1, "Total questions", BB); cell(mk, tr + 1, 2, f"=B{tr}+C{tr}", BB, center=True)
mk.freeze_panes = "E2"

# ---------- Release Calendar ----------
header(cal, 1, ["Release date", "Day", "Test ID", "Type", "Title", "Questions", "Marks", "Minutes", "Files ready for team check by"],
       [12, 6, 8, 16, 52, 9, 7, 8, 16])
import datetime as dt
for r, (d, tid, typ, title, q, mks, mins) in enumerate(rel, 2):
    dd = dt.date.fromisoformat(d)
    c = cell(cal, r, 1, dd, center=True); c.number_format = "DD-MMM-YYYY"
    cell(cal, r, 2, f"=TEXT(A{r},\"ddd\")", center=True)
    cell(cal, r, 3, tid, BB, center=True); cell(cal, r, 4, typ); cell(cal, r, 5, title)
    cell(cal, r, 6, q, center=True); cell(cal, r, 7, mks, center=True); cell(cal, r, 8, mins, center=True)
    c = cell(cal, r, 9, f"=A{r}-5", center=True); c.number_format = "DD-MMM-YYYY"
cal.freeze_panes = "A2"
cl = len(rel) + 1

# ---------- Overview ----------
ov.column_dimensions["A"].width = 46; ov.column_dimensions["B"].width = 16; ov.column_dimensions["C"].width = 16
ov.column_dimensions["D"].width = 16; ov.column_dimensions["E"].width = 16
ov["A1"] = "GATE ME 2027 Test Series — Coverage Plan"; ov["A1"].font = Font(name=F, bold=True, size=14)
ov["A2"] = "Every topic is tested in its topic test, then revised in the subject test and in full-length mocks."; ov["A2"].font = B
header_row = 4
for i, v in enumerate(["Test type", "Tests", "Questions each", "Total questions", "Marks each"], 1):
    c = ov.cell(header_row, i, v); c.font = H; c.fill = HF; c.border = BOX
rows = [("Topic tests", f"=COUNTIF('Release Calendar'!D2:D{cl},\"Topic test\")", 33, 50),
        ("Full-subject tests", f"=COUNTIF('Release Calendar'!D2:D{cl},\"Subject test\")", 33, 50),
        ("Full-length mocks (GATE pattern, 3 hr)", f"=COUNTIF('Release Calendar'!D2:D{cl},\"Full-length mock\")", 65, 100)]
for k, (a, b, c_, d) in enumerate(rows, header_row + 1):
    cell(ov, k, 1, a); cell(ov, k, 2, b, center=True); cell(ov, k, 3, c_, center=True)
    cell(ov, k, 4, f"=B{k}*C{k}", center=True); cell(ov, k, 5, d, center=True)
k = header_row + 4
cell(ov, k, 1, "TOTAL", BB); cell(ov, k, 2, f"=SUM(B{header_row+1}:B{k-1})", BB, center=True)
cell(ov, k, 4, f"=SUM(D{header_row+1}:D{k-1})", BB, center=True)

k += 2
for i, v in enumerate(["Revision check", "A topics", "B topics", "C topics", "All topics"], 1):
    c = ov.cell(k, i, v); c.font = H; c.fill = HF; c.border = BOX
R = f"'Topic Register'!"
lines = [("Number of topics", "COUNTIF({r}E2:E{l},\"{i}\")", f"COUNTA({R}A2:A{last})"),
         ("Fewest mock appearances", "_xlfn.MINIFS({r}W2:W{l},{r}E2:E{l},\"{i}\")", f"MIN({R}W2:W{last})"),
         ("Most mock appearances", "_xlfn.MAXIFS({r}W2:W{l},{r}E2:E{l},\"{i}\")", f"MAX({R}W2:W{last})"),
         ("Fewest revisions after the topic test", "_xlfn.MINIFS({r}X2:X{l},{r}E2:E{l},\"{i}\")", f"MIN({R}X2:X{last})"),
         ("Topics failing the revision rule", "COUNTIFS({r}E2:E{l},\"{i}\",{r}Y2:Y{l},\"NO\")", f"COUNTIF({R}Y2:Y{last},\"NO\")")]
for j, (lab, f, fall) in enumerate(lines, k + 1):
    cell(ov, j, 1, lab)
    for c, imp in enumerate("ABC", 2):
        cell(ov, j, c, "=" + f.format(r=R, l=last, i=imp), center=True)
    cell(ov, j, 5, "=" + fall, BB, center=True)
j += 2
notes = ["Importance: A = asked almost every year (≥4 mocks), B = regular (≥2 mocks), C = occasional (≥1 mock).",
         "Revisions after the topic test = full-subject test + every mock the topic appears in. Rule: at least 2 for every topic.",
         "Mock pattern: 65 Q / 100 marks / 180 min — 30 one-mark + 35 two-mark; GA 15 marks, Maths 13, core 72.",
         "Question types (as in GATE90): MCQ (−1/3, −2/3), integer and fill_ups (0 negative). Real GATE MSQ is not supported by the site.",
         "Mock dates: Sundays from 6 Dec, plus Wednesdays from 23 Dec. 'Files ready' = 5 days before release for the team's check."]
for n_, t_ in enumerate(notes, j):
    ov.cell(n_, 1, t_).font = B

wb.save("GATE_ME_2027_Test_Series_Coverage_Plan.xlsx")
print("saved")
